import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timedelta
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Sum, Count, Q, F
from decimal import Decimal

from accounts.decorators import admin_required
from core.models import Business
from services.models import ServiceJob, JobPart
from inventory.models import InventoryItem, Brand
from accounts.models import StaffProfile

def get_date_range(filter_type, start_date_str=None, end_date_str=None):
    now = timezone.now()
    today = timezone.localdate()

    if filter_type == 'today':
        start_date = today
        end_date = today
    elif filter_type == 'week':
        start_date = today - timedelta(days=7)
        end_date = today
    elif filter_type == 'month':
        start_date = today.replace(day=1)
        end_date = today
    elif filter_type == 'custom' and start_date_str and end_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        except ValueError:
            start_date = today.replace(day=1)
            end_date = today
    else:
        # Default to current month
        start_date = today.replace(day=1)
        end_date = today

    return start_date, end_date

@admin_required
def reports_index_view(request):
    workshop = Business.objects.first()
    filter_type = request.GET.get('range', 'month')
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    start_date, end_date = get_date_range(filter_type, start_date_str, end_date_str)

    jobs = ServiceJob.objects.filter(
        business=workshop,
        created_at__date__gte=start_date,
        created_at__date__lte=end_date
    )

    total_jobs = jobs.count()
    completed_jobs = jobs.filter(status__in=['COMPLETED', 'CLOSED']).count()
    pending_jobs = total_jobs - completed_jobs

    total_revenue = jobs.filter(payment_status='PAID').aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
    total_service_charges = jobs.filter(payment_status='PAID').aggregate(total=Sum('service_charge'))['total'] or Decimal('0.00')
    total_parts_charges = jobs.filter(payment_status='PAID').aggregate(total=Sum('parts_charge'))['total'] or Decimal('0.00')

    # Status distribution
    status_counts = jobs.values('status').annotate(count=Count('id'))
    completed_closed_count = jobs.filter(status__in=['COMPLETED', 'CLOSED']).count()
    in_progress_count = jobs.filter(status='IN_PROGRESS').count()
    waiting_parts_count = jobs.filter(status='WAITING_PARTS').count()
    assigned_new_count = jobs.filter(status__in=['ASSIGNED', 'NEW']).count()

    # Technician performance
    tech_stats = StaffProfile.objects.filter(business=workshop, role='TECHNICIAN').annotate(
        assigned_count=Count('user__assigned_jobs', filter=Q(user__assigned_jobs__created_at__date__gte=start_date, user__assigned_jobs__created_at__date__lte=end_date)),
        completed_count=Count('user__assigned_jobs', filter=Q(user__assigned_jobs__status__in=['COMPLETED', 'CLOSED'], user__assigned_jobs__created_at__date__gte=start_date, user__assigned_jobs__created_at__date__lte=end_date)),
    )

    # Brand distribution
    brand_stats = jobs.values('ac_brand_name').annotate(
        count=Count('id'),
        revenue=Sum('total_amount')
    ).order_by('-revenue')

    # Top Parts Consumed in this period
    top_parts = JobPart.objects.filter(
        job__created_at__date__gte=start_date,
        job__created_at__date__lte=end_date
    ).values('inventory_item__part_name', 'inventory_item__unit').annotate(
        total_qty=Sum('quantity'),
        total_amount=Sum('subtotal')
    ).order_by('-total_qty')[:10]

    # Low Stock Items
    low_stock_items = InventoryItem.objects.filter(
        business=workshop,
        quantity__lte=F('minimum_stock')
    ).select_related('brand')

    context = {
        'filter_type': filter_type,
        'start_date': start_date,
        'end_date': end_date,
        'total_jobs': total_jobs,
        'completed_jobs': completed_jobs,
        'pending_jobs': pending_jobs,
        'total_revenue': total_revenue,
        'total_service_charges': total_service_charges,
        'total_parts_charges': total_parts_charges,
        'status_counts': status_counts,
        'completed_closed_count': completed_closed_count,
        'in_progress_count': in_progress_count,
        'waiting_parts_count': waiting_parts_count,
        'assigned_new_count': assigned_new_count,
        'tech_stats': tech_stats,
        'brand_stats': brand_stats,
        'top_parts': top_parts,
        'low_stock_items': low_stock_items,
    }
    return render(request, 'reports/reports_index.html', context)

@admin_required
def export_jobs_csv(request):
    workshop = Business.objects.first()
    filter_type = request.GET.get('range', 'month')
    start_date, end_date = get_date_range(filter_type, request.GET.get('start_date'), request.GET.get('end_date'))

    jobs = ServiceJob.objects.filter(
        business=workshop,
        created_at__date__gte=start_date,
        created_at__date__lte=end_date
    ).select_related('customer', 'technician')

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="Service_Jobs_{start_date}_to_{end_date}.csv"'

    writer = csv.writer(response)
    writer.writerow([
        'Job ID', 'Date', 'Customer Name', 'Phone', 'Address',
        'AC Brand', 'AC Type', 'Service Type', 'Complaint',
        'Technician', 'Status', 'Service Charge (Rs.)',
        'Parts Charge (Rs.)', 'Discount (Rs.)', 'Total (Rs.)',
        'Payment Status', 'Payment Method'
    ])

    for j in jobs:
        writer.writerow([
            j.job_number,
            j.created_at.strftime('%Y-%m-%d %H:%M'),
            j.customer.name,
            j.customer.phone,
            j.customer.address,
            j.ac_brand_name,
            j.ac_type_name,
            j.service_type,
            j.complaint.replace('\n', ' '),
            j.technician.get_full_name() if j.technician else 'Unassigned',
            j.get_status_display(),
            float(j.service_charge),
            float(j.parts_charge),
            float(j.discount),
            float(j.total_amount),
            j.get_payment_status_display(),
            j.get_payment_method_display() or 'N/A'
        ])

    return response

@admin_required
def export_jobs_excel(request):
    workshop = Business.objects.first()
    filter_type = request.GET.get('range', 'month')
    start_date, end_date = get_date_range(filter_type, request.GET.get('start_date'), request.GET.get('end_date'))

    jobs = ServiceJob.objects.filter(
        business=workshop,
        created_at__date__gte=start_date,
        created_at__date__lte=end_date
    ).select_related('customer', 'technician')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Service Jobs Report"

    # Styling
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="475569")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title rows
    ws.merge_cells('A1:L1')
    ws['A1'] = f"{workshop.name} - Service Jobs Report"
    ws['A1'].font = title_font

    ws.merge_cells('A2:L2')
    ws['A2'] = f"Report Period: {start_date} to {end_date} | Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}"
    ws['A2'].font = subtitle_font

    headers = [
        'Job ID', 'Date', 'Customer Name', 'Phone', 'Address',
        'AC Brand', 'AC Type', 'Service Type', 'Technician',
        'Status', 'Service (Rs.)', 'Parts (Rs.)', 'Total (Rs.)', 'Payment'
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, j in enumerate(jobs, 5):
        row_values = [
            j.job_number,
            j.created_at.strftime('%Y-%m-%d %H:%M'),
            j.customer.name,
            j.customer.phone,
            j.customer.address,
            j.ac_brand_name,
            j.ac_type_name,
            j.service_type,
            j.technician.get_full_name() if j.technician else 'Unassigned',
            j.get_status_display(),
            float(j.service_charge),
            float(j.parts_charge),
            float(j.total_amount),
            f"{j.get_payment_status_display()} ({j.get_payment_method_display()})"
        ]
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_thin
            if col_idx in [11, 12, 13]:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Service_Jobs_{start_date}_to_{end_date}.xlsx"'
    return response

@admin_required
def export_inventory_excel(request):
    workshop = Business.objects.first()
    items = InventoryItem.objects.filter(business=workshop).select_related('brand', 'category').order_by('part_name')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory Status"

    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    low_stock_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    ws.merge_cells('A1:I1')
    ws['A1'] = f"{workshop.name} - Spare Parts Inventory Report"
    ws['A1'].font = title_font

    headers = [
        'Part Name', 'Code', 'Category', 'Brand / Type',
        'Current Stock', 'Min Stock', 'Unit', 'Purchase Price (Rs.)',
        'Selling Price (Rs.)', 'Stock Status'
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, itm in enumerate(items, 4):
        is_low = itm.is_low_stock
        row_values = [
            itm.part_name,
            itm.part_code or '-',
            itm.category.name if itm.category else '-',
            f"{itm.brand.name} ({itm.get_inventory_type_display()})" if itm.brand else it.get_inventory_type_display(),
            float(itm.quantity),
            float(itm.minimum_stock),
            itm.unit,
            float(itm.purchase_price),
            float(itm.selling_price),
            "LOW STOCK ALERT" if is_low else "Normal Stock"
        ]
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_thin
            if is_low and col_idx == 10:
                cell.fill = low_stock_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Inventory_Status_{timezone.localdate()}.xlsx"'
    return response
