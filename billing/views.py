from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Q, Sum, Count
from django.utils import timezone
from decimal import Decimal
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from accounts.decorators import admin_required
from core.models import Business, ActivityLog
from inventory.models import Brand
from .models import Invoice, Payment, Dealer
from services.models import ServiceJob

@login_required
def invoice_list_view(request):
    workshop = Business.objects.first()
    search = request.GET.get('search', '').strip()
    status_filter = request.GET.get('status', '')
    tech_id = request.GET.get('technician', '')
    date_filter = request.GET.get('date', '')
    from_date = request.GET.get('from_date', '').strip()
    to_date = request.GET.get('to_date', '').strip()

    invoices = Invoice.objects.filter(business=workshop).select_related('customer', 'job', 'technician', 'dealer')

    if search:
        invoices = invoices.filter(
            Q(invoice_number__icontains=search) |
            Q(job__job_number__icontains=search) |
            Q(customer__name__icontains=search) |
            Q(customer__phone__icontains=search) |
            Q(customer__address__icontains=search)
        )
    if status_filter:
        invoices = invoices.filter(payment_status=status_filter)
    if tech_id:
        invoices = invoices.filter(technician_id=tech_id)

    today = timezone.localdate()
    if date_filter == 'today':
        invoices = invoices.filter(created_at__date=today)
    elif date_filter == 'week':
        start_week = today - timezone.timedelta(days=7)
        invoices = invoices.filter(created_at__date__gte=start_week, created_at__date__lte=today)
    elif date_filter == 'month':
        start_month = today.replace(day=1)
        invoices = invoices.filter(created_at__date__gte=start_month, created_at__date__lte=today)

    if from_date:
        invoices = invoices.filter(created_at__date__gte=from_date)
    if to_date:
        invoices = invoices.filter(created_at__date__lte=to_date)

    invoices = invoices.order_by('-created_at')

    # Aggregates
    total_invoiced = invoices.filter(is_warranty=False).aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
    total_paid = invoices.filter(payment_status='PAID').aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
    total_due = invoices.filter(payment_status='PENDING', is_warranty=False).aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
    total_warranty_claims = invoices.filter(is_warranty=True).aggregate(total=Sum('dealer_claim_amount'))['total'] or Decimal('0.00')

    from accounts.models import StaffProfile
    technicians = StaffProfile.objects.filter(business=workshop, role='TECHNICIAN').select_related('user')

    context = {
        'invoices': invoices,
        'search': search,
        'status_filter': status_filter,
        'selected_tech': tech_id,
        'date_filter': date_filter,
        'from_date': from_date,
        'to_date': to_date,
        'technicians': technicians,
        'total_invoiced': total_invoiced,
        'total_paid': total_paid,
        'total_due': total_due,
        'total_warranty_claims': total_warranty_claims,
    }
    return render(request, 'billing/invoice_list.html', context)

@admin_required
def invoice_detail_view(request, pk):
    invoice = get_object_or_404(
        Invoice.objects.select_related('business', 'customer', 'job', 'technician', 'dealer'),
        pk=pk
    )
    job = invoice.job
    parts = job.parts_used.select_related('inventory_item').all()
    work_logs = job.work_logs.all()
    payments = invoice.payments.all()

    context = {
        'invoice': invoice,
        'job': job,
        'parts': parts,
        'work_logs': work_logs,
        'payments': payments,
    }
    return render(request, 'billing/invoice_detail.html', context)

@admin_required
def invoice_record_payment_view(request, pk):
    invoice = get_object_or_404(Invoice, pk=pk)
    if request.method == 'POST':
        method = request.POST.get('payment_method', 'CASH')
        reference = request.POST.get('reference', '').strip()
        notes = request.POST.get('notes', '').strip()
        amount_paid_val = Decimal(request.POST.get('amount', str(invoice.total_amount)) or '0')

        invoice.amount_paid += amount_paid_val
        if invoice.amount_paid >= invoice.total_amount:
            invoice.payment_status = 'PAID'
        invoice.payment_method = method
        invoice.paid_at = timezone.now()
        invoice.save()

        # Sync to job as well
        job = invoice.job
        if invoice.payment_status == 'PAID':
            job.payment_status = 'PAID'
            if job.status == 'COMPLETED':
                job.status = 'CLOSED'
                job.closed_at = invoice.paid_at
        job.payment_method = method
        job.payment_reference = reference
        job.paid_at = invoice.paid_at
        job.save()

        Payment.objects.create(
            invoice=invoice,
            amount=amount_paid_val,
            payment_method=method,
            reference=reference,
            received_by=request.user,
            notes=notes
        )

        ActivityLog.objects.create(
            business=invoice.business,
            user=request.user,
            action=f"Payment recorded for {invoice.invoice_number} (Rs. {amount_paid_val})",
            action_type="PAYMENT",
            job_reference=job.job_number
        )

        messages.success(request, f"Payment of Rs. {amount_paid_val} recorded for {invoice.invoice_number}.")

    # Redirect back to referring page (e.g. credit ledger or invoice detail)
    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER')
    if next_url:
        return redirect(next_url)
    return redirect('billing:invoice_detail', pk=pk)

@admin_required
def credit_ledger_view(request):
    """
    Dedicated Due / Udharo Khata (Credit Ledger).
    Lists all customers with pending balances, contact numbers, addresses, and days overdue.
    """
    workshop = Business.objects.first()
    search = request.GET.get('search', '').strip()

    due_invoices = Invoice.objects.filter(
        business=workshop,
        is_warranty=False,
        payment_status='PENDING'
    ).select_related('customer', 'job', 'technician').order_by('-created_at')

    if search:
        due_invoices = due_invoices.filter(
            Q(customer__name__icontains=search) |
            Q(customer__phone__icontains=search) |
            Q(customer__address__icontains=search) |
            Q(invoice_number__icontains=search) |
            Q(job__job_number__icontains=search)
        )

    total_due_amount = due_invoices.aggregate(total=Sum('total_amount'))['total'] or Decimal('0.00')
    total_customers_count = due_invoices.values('customer').distinct().count()

    context = {
        'invoices': due_invoices,
        'search': search,
        'total_due_amount': total_due_amount,
        'total_customers_count': total_customers_count,
    }
    return render(request, 'billing/credit_ledger.html', context)

@login_required
def warranty_claims_view(request):
    """
    Dedicated Dealer Warranty Claims Hub (WTY).
    Tracks claims to Authorized Brand Dealers (LG, Samsung, GEM, CG, etc.).
    """
    workshop = Business.objects.first()
    dealer_id = request.GET.get('dealer', '')
    status_filter = request.GET.get('status', '')
    search = request.GET.get('search', '').strip()
    date_filter = request.GET.get('date', '')
    from_date = request.GET.get('from_date', '').strip()
    to_date = request.GET.get('to_date', '').strip()

    # Ensure warranty invoices are recorded in company warranty claims
    unsynced_wty_jobs = ServiceJob.objects.filter(business=workshop, is_warranty=True, invoice__isnull=True)
    for j in unsynced_wty_jobs:
        try:
            Invoice.sync_from_job(j)
        except Exception:
            pass

    claims = Invoice.objects.filter(
        business=workshop,
        is_warranty=True
    ).select_related('customer', 'job', 'dealer', 'technician').order_by('-created_at')

    if dealer_id:
        claims = claims.filter(dealer_id=dealer_id)
    if status_filter:
        claims = claims.filter(dealer_claim_status=status_filter)
    if search:
        claims = claims.filter(
            Q(dealer__name__icontains=search) |
            Q(customer__name__icontains=search) |
            Q(customer__phone__icontains=search) |
            Q(job__job_number__icontains=search) |
            Q(invoice_number__icontains=search)
        )

    today = timezone.localdate()
    if date_filter == 'today':
        claims = claims.filter(created_at__date=today)
    elif date_filter == 'week':
        start_week = today - timezone.timedelta(days=7)
        claims = claims.filter(created_at__date__gte=start_week, created_at__date__lte=today)
    elif date_filter == 'month':
        start_month = today.replace(day=1)
        claims = claims.filter(created_at__date__gte=start_month, created_at__date__lte=today)

    if from_date:
        claims = claims.filter(created_at__date__gte=from_date)
    if to_date:
        claims = claims.filter(created_at__date__lte=to_date)

    total_claim_amount = claims.aggregate(total=Sum('dealer_claim_amount'))['total'] or Decimal('0.00')
    settled_claim_amount = claims.filter(dealer_claim_status='SETTLED').aggregate(total=Sum('dealer_claim_amount'))['total'] or Decimal('0.00')
    pending_claim_amount = claims.filter(dealer_claim_status__in=['PENDING', 'SUBMITTED']).aggregate(total=Sum('dealer_claim_amount'))['total'] or Decimal('0.00')

    dealers = Dealer.objects.filter(business=workshop, is_active=True)

    context = {
        'claims': claims,
        'dealers': dealers,
        'selected_dealer': dealer_id,
        'selected_status': status_filter,
        'search': search,
        'date_filter': date_filter,
        'from_date': from_date,
        'to_date': to_date,
        'total_claim_amount': total_claim_amount,
        'settled_claim_amount': settled_claim_amount,
        'pending_claim_amount': pending_claim_amount,
    }
    return render(request, 'billing/warranty_claims.html', context)

@login_required
def warranty_claims_export_excel(request):
    workshop = Business.objects.first()
    dealer_id = request.GET.get('dealer', '')
    status_filter = request.GET.get('status', '')
    search = request.GET.get('search', '').strip()
    date_filter = request.GET.get('date', '')
    from_date = request.GET.get('from_date', '').strip()
    to_date = request.GET.get('to_date', '').strip()

    # Ensure all warranty jobs have synced invoices
    unsynced_wty_jobs = ServiceJob.objects.filter(business=workshop, is_warranty=True, invoice__isnull=True)
    for j in unsynced_wty_jobs:
        try:
            Invoice.sync_from_job(j)
        except Exception:
            pass

    claims = Invoice.objects.filter(
        business=workshop,
        is_warranty=True
    ).select_related('customer', 'job', 'dealer', 'technician').order_by('-created_at')

    selected_dealer_name = "All Authorized Brand Companies"
    if dealer_id:
        claims = claims.filter(dealer_id=dealer_id)
        dealer_obj = Dealer.objects.filter(id=dealer_id).first()
        if dealer_obj:
            selected_dealer_name = dealer_obj.name

    if status_filter:
        claims = claims.filter(dealer_claim_status=status_filter)
    if search:
        claims = claims.filter(
            Q(dealer__name__icontains=search) |
            Q(customer__name__icontains=search) |
            Q(customer__phone__icontains=search) |
            Q(job__job_number__icontains=search) |
            Q(invoice_number__icontains=search)
        )

    today = timezone.localdate()
    if date_filter == 'today':
        claims = claims.filter(created_at__date=today)
    elif date_filter == 'week':
        start_week = today - timezone.timedelta(days=7)
        claims = claims.filter(created_at__date__gte=start_week, created_at__date__lte=today)
    elif date_filter == 'month':
        start_month = today.replace(day=1)
        claims = claims.filter(created_at__date__gte=start_month, created_at__date__lte=today)

    if from_date:
        claims = claims.filter(created_at__date__gte=from_date)
    if to_date:
        claims = claims.filter(created_at__date__lte=to_date)

    # Build Excel Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Warranty Claims"

    title_font = Font(name='Calibri', size=15, bold=True, color='1E3A8A')
    subtitle_font = Font(name='Calibri', size=10, italic=True, color='475569')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
    data_font = Font(name='Calibri', size=10, color='0F172A')
    bold_font = Font(name='Calibri', size=10, bold=True, color='0F172A')
    total_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    workshop_name = workshop.name if workshop else "AC & Appliance Service Workshop"
    ws.merge_cells('A1:L1')
    ws['A1'] = f"{workshop_name} - Company Warranty Claims Report"
    ws['A1'].font = title_font

    ws.merge_cells('A2:L2')
    ws['A2'] = f"Dealer / Company: {selected_dealer_name} | Generated: {timezone.now().strftime('%Y-%m-%d %I:%M %p')}"
    ws['A2'].font = subtitle_font

    start_row = 4
    if from_date or to_date or date_filter:
        ws.merge_cells('A3:L3')
        ws['A3'] = f"Filter: {date_filter.capitalize() if date_filter else ''} {f'From: {from_date}' if from_date else ''} {f'To: {to_date}' if to_date else ''}"
        ws['A3'].font = subtitle_font
        start_row = 5

    headers = [
        "S.N.",
        "Claim Voucher #",
        "Job ID",
        "Date (AD)",
        "Date (BS)",
        "Dealer / Company",
        "Customer Name",
        "Customer Phone",
        "Appliance & Brand",
        "Job Type",
        "Complaint / Problem",
        "Claim Amount (Rs.)",
        "Claim Status"
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center' if col_idx in [1, 4, 5, 13] else 'left', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[start_row].height = 24

    current_row = start_row + 1
    total_claim_amount = Decimal('0.00')

    for sn, claim in enumerate(claims, 1):
        job = claim.job
        dealer_name = claim.dealer.name if claim.dealer else (claim.company_claim_name or "General Dealer")
        ad_date = claim.created_at.strftime('%Y-%m-%d')
        bs_date = job.date_bs if job and hasattr(job, 'date_bs') else ""
        appliance_info = f"{job.get_appliance_type_display()} - {job.ac_brand_name or (job.brand.name if job.brand else '')}" if job else "-"
        job_type = job.get_job_type_display() if job else "-"
        complaint = job.complaint if job else "-"
        claim_amt = claim.dealer_claim_amount or (job.dealer_claim_amount if job else Decimal('0.00'))
        status_disp = claim.get_dealer_claim_status_display() if hasattr(claim, 'get_dealer_claim_status_display') else claim.dealer_claim_status

        total_claim_amount += claim_amt

        row_data = [
            sn,
            claim.invoice_number,
            job.job_number if job else "-",
            ad_date,
            bs_date,
            dealer_name,
            claim.customer.name if claim.customer else "-",
            claim.customer.phone if claim.customer else "-",
            appliance_info,
            job_type,
            complaint,
            float(claim_amt),
            status_disp
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = Alignment(horizontal='center')
            elif col_idx in [4, 5]:
                cell.alignment = Alignment(horizontal='center')
            elif col_idx == 12:
                cell.alignment = Alignment(horizontal='right')
                cell.number_format = '#,##0.00'
            elif col_idx == 13:
                cell.alignment = Alignment(horizontal='center')
                if status_disp == 'Settled':
                    cell.font = Font(name='Calibri', size=10, bold=True, color='15803D')
                elif status_disp in ['Submitted', 'Pending']:
                    cell.font = Font(name='Calibri', size=10, bold=True, color='D97706')

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # TOTAL ROW
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=11)
    tot_label = ws.cell(row=current_row, column=1, value="TOTAL WARRANTY CLAIMS (Rs.)")
    tot_label.font = bold_font
    tot_label.alignment = Alignment(horizontal='right', vertical='center')
    tot_label.fill = total_fill
    tot_label.border = thin_border

    tot_val = ws.cell(row=current_row, column=12, value=float(total_claim_amount))
    tot_val.font = Font(name='Calibri', size=11, bold=True, color='1E40AF')
    tot_val.alignment = Alignment(horizontal='right', vertical='center')
    tot_val.number_format = '#,##0.00'
    tot_val.fill = total_fill
    tot_val.border = thin_border

    tot_end = ws.cell(row=current_row, column=13, value=f"{claims.count()} claims")
    tot_end.font = bold_font
    tot_end.alignment = Alignment(horizontal='center', vertical='center')
    tot_end.fill = total_fill
    tot_end.border = thin_border
    ws.row_dimensions[current_row].height = 22

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row < start_row:
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    safe_dealer_slug = selected_dealer_name.replace(" ", "_").replace("/", "-")[:30]
    filename = f"Warranty_Claims_{safe_dealer_slug}_{today.strftime('%Y%m%d')}.xlsx"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

@admin_required
def warranty_claim_update_status(request, pk):
    invoice = get_object_or_404(Invoice, pk=pk, is_warranty=True)
    if request.method == 'POST':
        new_status = request.POST.get('dealer_claim_status', 'SUBMITTED')
        invoice.dealer_claim_status = new_status
        if new_status == 'SETTLED':
            invoice.paid_at = timezone.now()
            invoice.payment_status = 'PAID'
            invoice.job.payment_status = 'PAID'
            invoice.job.paid_at = invoice.paid_at
            if invoice.job.status == 'COMPLETED':
                invoice.job.status = 'CLOSED'
                invoice.job.closed_at = invoice.paid_at
            invoice.job.save()
        invoice.save()

        # Update job field as well
        invoice.job.dealer_claim_status = new_status
        invoice.job.save(update_fields=['dealer_claim_status'])

        ActivityLog.objects.create(
            business=invoice.business,
            user=request.user,
            action=f"Warranty claim {invoice.invoice_number} updated to {new_status}",
            action_type="PAYMENT",
            job_reference=invoice.job.job_number,
            details=f"Dealer: {invoice.dealer.name if invoice.dealer else 'General'}"
        )

        messages.success(request, f"Claim {invoice.invoice_number} status updated to {new_status}.")

    return redirect('billing:warranty_claims')

@admin_required
def dealers_list_view(request):
    workshop = Business.objects.first()
    dealers = Dealer.objects.filter(business=workshop).select_related('brand').order_by('name')
    brands = Brand.objects.filter(business=workshop)

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        brand_id = request.POST.get('brand')
        contact_person = request.POST.get('contact_person', '').strip()
        phone = request.POST.get('phone', '').strip()
        address = request.POST.get('address', '').strip()
        notes = request.POST.get('notes', '').strip()

        if name:
            brand_obj = Brand.objects.filter(id=brand_id).first() if brand_id else None
            Dealer.objects.create(
                business=workshop,
                name=name,
                brand=brand_obj,
                contact_person=contact_person,
                phone=phone,
                address=address or "Damak, Jhapa",
                notes=notes
            )
            messages.success(request, f"Dealer '{name}' registered successfully.")
            return redirect('billing:dealer_list')

    context = {
        'dealers': dealers,
        'brands': brands,
    }
    return render(request, 'billing/dealer_list.html', context)
